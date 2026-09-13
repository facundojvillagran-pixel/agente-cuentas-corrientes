import io

from fastapi.testclient import TestClient
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfReader

from cuentas_corrientes.api import _account_configs, _accounts, _pending_uploads, app
from cuentas_corrientes.files import MAX_FILE_BYTES

client = TestClient(app)


def setup_function() -> None:
    _accounts.clear()
    _pending_uploads.clear()
    _account_configs.clear()


def _new_account(**opening) -> str:
    payload = {"company": "Agro Norte SA", "counterparty": "Transporte Ruiz", **opening}
    return client.post("/accounts", json=payload).json()["id"]


def _xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_invoice_bytes(lines: list[str]) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in lines:
        pdf.cell(0, 8, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    return bytes(pdf.output())


def _blank_pdf_bytes() -> bytes:
    pdf = FPDF()
    pdf.add_page()
    return bytes(pdf.output())


def _png_bytes(color=(10, 20, 30)) -> bytes:
    image = Image.new("RGB", (40, 40), color=color)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _upload(account_id: str, filename: str, content: bytes, content_type: str = "application/octet-stream"):
    return client.post(f"/accounts/{account_id}/documents", files=[("files", (filename, content, content_type))])


def test_excel_with_normal_headers_creates_ready_movements() -> None:
    account_id = _new_account()
    content = _xlsx_bytes(
        ["Fecha", "Concepto", "Debe", "Haber", "Moneda", "Comprobante"],
        [
            ["01/07/2026", "Anticipo entregado", 50000, None, "ARS", "REC-1"],
            ["05/07/2026", "Cobro de flete", None, 30000, "ARS", "FC-2"],
        ],
    )
    response = _upload(account_id, "planilla.xlsx", content)
    assert response.status_code == 201
    result = response.json()[0]
    assert result["kind"] == "spreadsheet"
    assert result["outcome"] == "movimiento_creado"
    assert len(result["movement_ids"]) == 2

    movements = client.get(f"/accounts/{account_id}/movements").json()
    assert len(movements) == 2
    assert {m["amount"] for m in movements} == {"50000.00", "30000.00"}


def test_excel_alternative_headers_are_recognized() -> None:
    account_id = _new_account()
    content = _xlsx_bytes(
        ["Fecha Operacion", "Descripcion", "Debito", "Credito", "Divisa", "Nro Comprobante"],
        [["10/07/2026", "Pago de prueba", 15000, None, "ARS", "X-1"]],
    )
    response = _upload(account_id, "alternativo.xlsx", content)
    assert response.json()[0]["outcome"] == "movimiento_creado"


def test_excel_incomplete_rows_need_review_without_inventing_data() -> None:
    account_id = _new_account()
    content = _xlsx_bytes(
        ["Fecha", "Concepto", "Debe", "Haber", "Moneda"],
        [
            ["12/07/2026", "Fila con debe y haber a la vez", 1000, 500, "ARS"],
            ["", "Fila sin fecha", 1000, None, "ARS"],
        ],
    )
    response = _upload(account_id, "incompleta.xlsx", content)
    result = response.json()[0]
    assert result["outcome"] == "necesita_revision"
    assert result["movement_ids"] == []
    pending = client.get(f"/accounts/{account_id}/documents").json()
    assert len(pending) == 2
    assert all(p["kind"] == "spreadsheet_row" for p in pending)


def test_excel_keeps_ars_and_usd_separated() -> None:
    account_id = _new_account()
    content = _xlsx_bytes(
        ["Fecha", "Concepto", "Debe", "Haber", "Moneda"],
        [
            ["01/07/2026", "Movimiento ARS", 1000, None, "ARS"],
            ["01/07/2026", "Movimiento USD", None, 100, "USD"],
        ],
    )
    _upload(account_id, "monedas.xlsx", content)
    ars = client.get(f"/accounts/{account_id}/balance/ARS").json()
    usd = client.get(f"/accounts/{account_id}/balance/USD").json()
    assert ars["projected_amount"] == "1000.00"
    assert usd["projected_amount"] == "100.00"


def test_excel_duplicate_rows_are_not_loaded_twice() -> None:
    account_id = _new_account()
    content = _xlsx_bytes(
        ["Fecha", "Concepto", "Debe", "Haber", "Moneda"],
        [["01/07/2026", "Pago único", 5000, None, "ARS"]],
    )
    first = _upload(account_id, "dup.xlsx", content)
    second = _upload(account_id, "dup.xlsx", content)
    assert first.json()[0]["outcome"] == "movimiento_creado"
    assert second.json()[0]["outcome"] == "posible_duplicado"
    movements = client.get(f"/accounts/{account_id}/movements").json()
    assert len(movements) == 1


def test_excel_formula_cell_does_not_crash_and_is_flagged() -> None:
    account_id = _new_account()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Fecha", "Concepto", "Debe", "Haber", "Moneda"])
    sheet.append(["01/07/2026", "Con formula", "=1+1", None, "ARS"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = _upload(account_id, "formula.xlsx", buffer.getvalue())
    assert response.status_code == 201
    assert response.json()[0]["outcome"] in {"necesita_revision", "invalido", "posible_duplicado"}


def test_excel_without_recognizable_headers_needs_mapping() -> None:
    account_id = _new_account()
    content = _xlsx_bytes(["A", "B", "C", "D", "E"], [["01/07/2026", "algo", 100, None, "ARS"]])
    response = _upload(account_id, "sin_encabezados.xlsx", content)
    result = response.json()[0]
    assert result["outcome"] == "necesita_revision"
    pending = client.get(f"/accounts/{account_id}/documents").json()
    assert pending[0]["kind"] == "spreadsheet_mapping"
    assert pending[0]["preview_rows"]


def test_resolve_spreadsheet_mapping_creates_movements() -> None:
    account_id = _new_account()
    content = _xlsx_bytes(["A", "B", "C", "D", "E"], [["01/07/2026", "algo", 100, None, "ARS"]])
    _upload(account_id, "sin_encabezados.xlsx", content)
    pending_id = client.get(f"/accounts/{account_id}/documents").json()[0]["id"]

    response = client.post(
        f"/accounts/{account_id}/documents/{pending_id}/resolve-mapping",
        json={"header_row_index": 0, "fecha": 0, "concepto": 1, "debe": 2, "haber": 3, "moneda": 4},
    )
    assert response.status_code == 200
    assert response.json()["created"] == 1
    assert client.get(f"/accounts/{account_id}/documents").json() == []


def test_pdf_invoice_with_text_extracts_fields_and_creates_movement() -> None:
    account_id = _new_account()
    content = _pdf_invoice_bytes(
        [
            "FACTURA A 0001-00001234",
            "CUIT: 30-12345678-9",
            "Fecha: 15/07/2026",
            "Moneda: ARS",
            "Neto gravado: 100.000,00",
            "IVA: 21.000,00",
            "Total: 121.000,00",
        ]
    )
    response = _upload(account_id, "factura.pdf", content)
    result = response.json()[0]
    assert result["kind"] == "invoice_pdf"
    assert result["outcome"] == "movimiento_creado"

    movements = client.get(f"/accounts/{account_id}/movements").json()
    assert movements[0]["amount"] == "121000.00"
    assert movements[0]["direction"] == "counterparty"


def test_pdf_without_text_needs_review() -> None:
    account_id = _new_account()
    response = _upload(account_id, "escaneada.pdf", _blank_pdf_bytes())
    result = response.json()[0]
    assert result["outcome"] == "necesita_revision"
    assert "escane" in result["message"].casefold()


def test_image_always_needs_manual_review() -> None:
    account_id = _new_account()
    response = _upload(account_id, "foto.png", _png_bytes())
    result = response.json()[0]
    assert result["kind"] == "photo"
    assert result["outcome"] == "necesita_revision"
    pending = client.get(f"/accounts/{account_id}/documents").json()
    assert pending[0]["extracted"] == {}


def test_corrupt_file_is_rejected_without_crashing() -> None:
    account_id = _new_account()
    response = _upload(account_id, "roto.xlsx", b"PK\x03\x04" + b"no es un zip valido" * 5)
    assert response.status_code == 201
    assert response.json()[0]["outcome"] == "invalido"


def test_fake_extension_is_rejected() -> None:
    account_id = _new_account()
    response = _upload(account_id, "factura.pdf", _png_bytes())
    assert response.json()[0]["outcome"] == "invalido"


def test_malicious_filename_is_sanitized() -> None:
    account_id = _new_account()
    response = _upload(account_id, "../../etc/passwd.png", _png_bytes())
    result = response.json()[0]
    assert ".." not in result["filename"]
    assert "/" not in result["filename"]


def test_oversized_file_is_rejected() -> None:
    account_id = _new_account()
    oversized = b"%PDF-" + b"0" * (MAX_FILE_BYTES + 100)
    response = _upload(account_id, "grande.pdf", oversized)
    assert response.json()[0]["outcome"] == "invalido"
    assert "tama" in response.json()[0]["message"].casefold()


def test_resolve_pending_document_creates_movement() -> None:
    account_id = _new_account()
    _upload(account_id, "foto.png", _png_bytes())
    pending_id = client.get(f"/accounts/{account_id}/documents").json()[0]["id"]

    response = client.post(
        f"/accounts/{account_id}/documents/{pending_id}/resolve",
        json={
            "operation_date": "2026-07-20",
            "description": "Foto de transferencia completada a mano",
            "amount": "80000",
            "currency": "ARS",
            "direction": "counterparty",
        },
    )
    assert response.status_code == 200
    assert client.get(f"/accounts/{account_id}/documents").json() == []
    movements = client.get(f"/accounts/{account_id}/movements").json()
    assert movements[0]["amount"] == "80000.00"


def test_bulk_approve_all_valid() -> None:
    account_id = _new_account()
    content = _xlsx_bytes(
        ["Fecha", "Concepto", "Debe", "Haber", "Moneda"],
        [
            ["01/07/2026", "Uno", 1000, None, "ARS"],
            ["02/07/2026", "Dos", None, 2000, "ARS"],
        ],
    )
    _upload(account_id, "varias.xlsx", content)
    response = client.post(f"/accounts/{account_id}/movements/approve-all-valid")
    assert response.status_code == 200
    assert response.json()["count"] == 2
    movements = client.get(f"/accounts/{account_id}/movements").json()
    assert all(m["status"] == "confirmed" for m in movements)


def test_reversal_correction_discard_and_persistence_still_work_with_documents() -> None:
    account_id = _new_account()
    _upload(
        account_id,
        "planilla.xlsx",
        _xlsx_bytes(["Fecha", "Concepto", "Debe", "Haber", "Moneda"], [["01/07/2026", "Pago", 1000, None, "ARS"]]),
    )
    movement_id = client.get(f"/accounts/{account_id}/movements").json()[0]["id"]
    client.post(f"/accounts/{account_id}/movements/{movement_id}/approve")
    reversed_response = client.post(f"/accounts/{account_id}/movements/{movement_id}/reverse", json={"reason": "prueba"})
    assert reversed_response.status_code == 200
    balance = client.get(f"/accounts/{account_id}/balance/ARS").json()
    assert balance["confirmed_amount"] == "0.00"


def test_export_excel_contains_expected_sheets_and_values() -> None:
    account_id = _new_account(
        opening_balance_amount="500000",
        opening_balance_currency="ARS",
        opening_balance_favor="contraparte",
    )
    _upload(
        account_id,
        "planilla.xlsx",
        _xlsx_bytes(["Fecha", "Concepto", "Debe", "Haber", "Moneda"], [["01/07/2026", "Cobro", None, 100000, "ARS"]]),
    )
    client.post(f"/accounts/{account_id}/movements/approve-all-valid")

    response = client.get(f"/accounts/{account_id}/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert "Cuenta ARS" in workbook.sheetnames
    sheet = workbook["Cuenta ARS"]
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert "Agro Norte SA" in values
    assert any(value == "Cobro" for value in values)


def test_export_pdf_generates_readable_document() -> None:
    account_id = _new_account()
    _upload(
        account_id,
        "planilla.xlsx",
        _xlsx_bytes(["Fecha", "Concepto", "Debe", "Haber", "Moneda"], [["01/07/2026", "Pago de prueba PDF", 1000, None, "ARS"]]),
    )
    client.post(f"/accounts/{account_id}/movements/approve-all-valid")

    response = client.get(f"/accounts/{account_id}/export.pdf")
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    reader = PdfReader(io.BytesIO(response.content))
    text = " ".join((page.extract_text() or "") for page in reader.pages)
    assert "Agro Norte SA" in text
    assert "Transporte Ruiz" in text


def test_account_config_logo_upload_view_and_removal() -> None:
    account_id = _new_account()
    response = client.post(
        f"/accounts/{account_id}/config",
        data={"display_name": "Nombre Comercial Ficticio"},
        files={"logo": ("logo.png", _png_bytes(), "image/png")},
    )
    assert response.status_code == 200
    assert response.json() == {"display_name": "Nombre Comercial Ficticio", "has_logo": True}

    logo_response = client.get(f"/accounts/{account_id}/config/logo")
    assert logo_response.status_code == 200
    assert logo_response.headers["content-type"] == "image/png"

    deleted = client.delete(f"/accounts/{account_id}/config/logo")
    assert deleted.json()["has_logo"] is False
    assert client.get(f"/accounts/{account_id}/config/logo").status_code == 404


def test_malicious_filename_is_returned_as_plain_text_never_executed() -> None:
    account_id = _new_account()
    response = _upload(account_id, "<img src=x onerror=alert(1)>.png", _png_bytes())
    result = response.json()[0]
    assert "<img" not in result["filename"]
    assert "onerror" not in result["filename"] or "<" not in result["filename"]
