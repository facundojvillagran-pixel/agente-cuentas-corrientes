from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from fpdf import FPDF
from fpdf.enums import XPos, YPos

from .documents import parse_vat_note
from .domain import Account, MovementStatus

_MONEY_FORMAT = "#,##0.00"


@dataclass
class ReportRow:
    operation_date: date
    description: str
    reference: str
    debe: Decimal | None
    haber: Decimal | None
    running_balance: Decimal
    status_note: str
    vat: Decimal | None


@dataclass
class CurrencyReport:
    currency: str
    opening_balance: Decimal
    confirmed_rows: list[ReportRow]
    pending_rows: list[ReportRow]
    confirmed_balance: Decimal
    projected_balance: Decimal
    unassigned_total: Decimal


def _build_rows(account: Account, currency: str, statuses: set[MovementStatus], start_balance: Decimal) -> tuple[list[ReportRow], Decimal]:
    reversed_ids = {m.reversal_of for m in account.movements if m.reversal_of}
    rows: list[ReportRow] = []
    running = start_balance
    for movement in account.ordered_movements():
        if movement.currency != currency or movement.is_opening_balance or movement.status not in statuses:
            continue
        running += movement.signed_amount
        debe = movement.amount if movement.direction.value == "company" else None
        haber = movement.amount if movement.direction.value == "counterparty" else None
        status_note = ""
        if movement.reversal_of:
            status_note = "Reversión"
        elif movement.id in reversed_ids:
            status_note = "Revertido"
        vat = parse_vat_note(movement.evidence[0].reference if movement.evidence else None)
        rows.append(ReportRow(movement.operation_date, movement.description, movement.external_reference or "", debe, haber, running, status_note, vat))
    return rows, running


def prepare_currency_report(account: Account, currency: str) -> CurrencyReport:
    currency = currency.upper()
    opening = account.opening_balances().get(currency, Decimal("0.00"))
    confirmed_rows, confirmed_balance = _build_rows(account, currency, {MovementStatus.CONFIRMED}, opening)
    pending_rows, _ = _build_rows(account, currency, {MovementStatus.REVIEW}, confirmed_balance)
    return CurrencyReport(
        currency=currency,
        opening_balance=opening,
        confirmed_rows=confirmed_rows,
        pending_rows=pending_rows,
        confirmed_balance=confirmed_balance,
        projected_balance=account.balances(projected=True).get(currency, Decimal("0.00")),
        unassigned_total=account.unassigned_total(currency),
    )


def currencies_in_use(account: Account) -> list[str]:
    found = {m.currency for m in account.movements}
    return sorted(found) or ["ARS"]


def build_excel_report(account: Account, *, display_name: str, logo_bytes: bytes | None) -> bytes:
    workbook = Workbook()
    first_sheet = True
    header_font = Font(bold=True)
    for currency in currencies_in_use(account):
        report = prepare_currency_report(account, currency)
        sheet = workbook.active if first_sheet else workbook.create_sheet()
        first_sheet = False
        sheet.title = f"Cuenta {currency}"[:31]

        row_cursor = 1
        sheet.cell(row=row_cursor, column=1, value=display_name).font = Font(bold=True, size=14)
        row_cursor += 1
        sheet.cell(row=row_cursor, column=1, value=f"Empresa: {account.company}    Contraparte: {account.counterparty}")
        row_cursor += 1
        sheet.cell(row=row_cursor, column=1, value=f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        row_cursor += 1
        sheet.cell(row=row_cursor, column=1, value=f"Moneda: {currency}")
        row_cursor += 2

        sheet.cell(row=row_cursor, column=1, value="Saldo inicial")
        sheet.cell(row=row_cursor, column=2, value=float(report.opening_balance)).number_format = _MONEY_FORMAT
        row_cursor += 2

        headers = ["Fecha", "Concepto", "Comprobante", "Debe", "Haber", "Saldo acumulado", "IVA", "Estado"]
        header_row = row_cursor
        for col, title in enumerate(headers, start=1):
            cell = sheet.cell(row=header_row, column=col, value=title)
            cell.font = header_font
        row_cursor += 1

        for row in report.confirmed_rows:
            sheet.cell(row=row_cursor, column=1, value=row.operation_date.isoformat())
            sheet.cell(row=row_cursor, column=2, value=row.description)
            sheet.cell(row=row_cursor, column=3, value=row.reference)
            if row.debe is not None:
                sheet.cell(row=row_cursor, column=4, value=float(row.debe)).number_format = _MONEY_FORMAT
            if row.haber is not None:
                sheet.cell(row=row_cursor, column=5, value=float(row.haber)).number_format = _MONEY_FORMAT
            sheet.cell(row=row_cursor, column=6, value=float(row.running_balance)).number_format = _MONEY_FORMAT
            if row.vat is not None:
                sheet.cell(row=row_cursor, column=7, value=float(row.vat)).number_format = _MONEY_FORMAT
            sheet.cell(row=row_cursor, column=8, value=row.status_note)
            row_cursor += 1

        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1).coordinate
        for col in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 22

        row_cursor += 1
        sheet.cell(row=row_cursor, column=1, value="Anticipos sin asignar")
        sheet.cell(row=row_cursor, column=2, value=float(report.unassigned_total)).number_format = _MONEY_FORMAT
        row_cursor += 1
        sheet.cell(row=row_cursor, column=1, value="Saldo final confirmado").font = header_font
        sheet.cell(row=row_cursor, column=2, value=float(report.confirmed_balance)).number_format = _MONEY_FORMAT
        row_cursor += 1
        sheet.cell(row=row_cursor, column=1, value="Saldo proyectado (incluye pendientes)")
        sheet.cell(row=row_cursor, column=2, value=float(report.projected_balance)).number_format = _MONEY_FORMAT

        if report.pending_rows:
            pending_sheet = workbook.create_sheet(f"Pendientes {currency}"[:31])
            for col, title in enumerate(["Fecha", "Concepto", "Comprobante", "Debe", "Haber", "IVA"], start=1):
                pending_sheet.cell(row=1, column=col, value=title).font = header_font
            for index, row in enumerate(report.pending_rows, start=2):
                pending_sheet.cell(row=index, column=1, value=row.operation_date.isoformat())
                pending_sheet.cell(row=index, column=2, value=row.description)
                pending_sheet.cell(row=index, column=3, value=row.reference)
                if row.debe is not None:
                    pending_sheet.cell(row=index, column=4, value=float(row.debe)).number_format = _MONEY_FORMAT
                if row.haber is not None:
                    pending_sheet.cell(row=index, column=5, value=float(row.haber)).number_format = _MONEY_FORMAT
                if row.vat is not None:
                    pending_sheet.cell(row=index, column=6, value=float(row.vat)).number_format = _MONEY_FORMAT
            pending_sheet.freeze_panes = "A2"
            for col in range(1, 7):
                pending_sheet.column_dimensions[get_column_letter(col)].width = 22

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class _ReportPDF(FPDF):
    def __init__(self, *, display_name: str, subtitle: str, logo_bytes: bytes | None) -> None:
        super().__init__(orientation="L")
        self.display_name = display_name
        self.subtitle = subtitle
        self.logo_bytes = logo_bytes
        self.set_auto_page_break(auto=True, margin=15)

    def header(self) -> None:  # noqa: D102 (fpdf2 hook)
        if self.logo_bytes:
            try:
                self.image(io.BytesIO(self.logo_bytes), x=10, y=8, h=14)
            except Exception:
                pass
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 8, self.display_name, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
        self.set_font("Helvetica", "", 10)
        self.cell(0, 6, self.subtitle, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
        self.ln(2)

    def footer(self) -> None:  # noqa: D102 (fpdf2 hook)
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 8, f"Página {self.page_no()}", align="C")


def build_pdf_report(account: Account, *, display_name: str, logo_bytes: bytes | None) -> bytes:
    subtitle = f"Empresa: {account.company}  ·  Contraparte: {account.counterparty}  ·  Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    pdf = _ReportPDF(display_name=display_name, subtitle=subtitle, logo_bytes=logo_bytes)
    pdf.set_margins(10, 28, 10)

    headers = ["Fecha", "Concepto", "Comprobante", "Debe", "Haber", "Saldo acum.", "Estado"]
    widths = [22, 90, 40, 30, 30, 32, 30]

    for currency in currencies_in_use(account):
        report = prepare_currency_report(account, currency)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, f"Moneda: {currency}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 6, f"Saldo inicial: {report.opening_balance:,.2f} {currency}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(2)

        def draw_header() -> None:
            pdf.set_font("Helvetica", "B", 9)
            for width, title in zip(widths, headers):
                pdf.cell(width, 7, title, border=1)
            pdf.ln()
            pdf.set_font("Helvetica", "", 9)

        draw_header()
        for row in report.confirmed_rows:
            if pdf.get_y() > pdf.h - 25:
                pdf.add_page()
                draw_header()
            pdf.cell(widths[0], 6, row.operation_date.isoformat(), border=1)
            pdf.cell(widths[1], 6, row.description[:45], border=1)
            pdf.cell(widths[2], 6, row.reference[:20], border=1)
            pdf.cell(widths[3], 6, f"{row.debe:,.2f}" if row.debe else "", border=1, align="R")
            pdf.cell(widths[4], 6, f"{row.haber:,.2f}" if row.haber else "", border=1, align="R")
            pdf.cell(widths[5], 6, f"{row.running_balance:,.2f}", border=1, align="R")
            pdf.cell(widths[6], 6, row.status_note, border=1)
            pdf.ln()

        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 6, f"Anticipos sin asignar: {report.unassigned_total:,.2f} {currency}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.cell(0, 6, f"Saldo final confirmado: {report.confirmed_balance:,.2f} {currency}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.cell(0, 6, f"Saldo proyectado: {report.projected_balance:,.2f} {currency}", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if report.pending_rows:
            pdf.set_text_color(150, 90, 0)
            pdf.cell(0, 8, f"Atención: hay {len(report.pending_rows)} propuesta(s) pendiente(s) de aprobación en {currency}.", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_text_color(0, 0, 0)

    return bytes(pdf.output())
