from __future__ import annotations

import io
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

_HEADER_SYNONYMS: dict[str, tuple[str, ...]] = {
    "fecha": ("fecha", "fecha operacion", "fecha de operacion", "fecha operativa"),
    "concepto": ("concepto", "descripcion", "detalle", "concepto o descripcion"),
    "debe": ("debe", "debito", "cargo"),
    "haber": ("haber", "credito", "abono"),
    "saldo": ("saldo",),
    "moneda": ("moneda", "divisa"),
    "comprobante": ("comprobante", "factura", "nro comprobante", "numero de comprobante", "n comprobante"),
    "referencia": ("referencia", "ref"),
    "litros": ("litros", "lts", "lt", "cantidad de litros", "litros cargados"),
}
_REQUIRED_HEADERS = {"fecha", "concepto", "debe", "haber", "moneda"}
_MIN_MATCHED_HEADERS = 3
_MAX_HEADER_SEARCH_ROWS = 15
_DEFAULT_DESCRIPTION = "Movimiento importado de planilla"


def _normalize(text: object) -> str:
    if text is None:
        return ""
    normalized = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode("ascii")
    return normalized.strip().casefold()


def _match_header(cell_value: object) -> str | None:
    normalized = _normalize(cell_value)
    if not normalized:
        return None
    for field_name, synonyms in _HEADER_SYNONYMS.items():
        if normalized in synonyms:
            return field_name
    return None


@dataclass
class SheetPreview:
    sheet_name: str
    rows: list[list[str]]


@dataclass
class ReadyRow:
    row_number: int
    operation_date: date
    description: str
    amount: Decimal
    currency: str
    direction: str  # "company" | "counterparty"
    external_reference: str | None
    row_hash_seed: str


@dataclass
class AmbiguousRow:
    row_number: int
    reason: str
    extracted: dict[str, str] = field(default_factory=dict)


@dataclass
class FuelChargeRow:
    row_number: int
    operation_date: date
    litres: Decimal
    description: str
    currency: str
    external_reference: str | None
    row_hash_seed: str


@dataclass
class SpreadsheetIngestResult:
    sheet_names: list[str]
    sheet_name: str
    header_found: bool
    header_row_index: int | None
    ready_rows: list[ReadyRow]
    ambiguous_rows: list[AmbiguousRow]
    preview: SheetPreview
    fuel_charge_rows: list[FuelChargeRow] = field(default_factory=list)


def _parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for pattern in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _parse_amount(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            amount = Decimal(str(value))
        except InvalidOperation:
            return None
    else:
        text = str(value).strip().replace(".", "").replace(",", ".")
        try:
            amount = Decimal(text)
        except InvalidOperation:
            return None
    return amount if amount > 0 else None


def _find_header_row(rows: list[tuple[object, ...]]) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(rows[:_MAX_HEADER_SEARCH_ROWS]):
        mapping: dict[str, int] = {}
        for col_index, cell in enumerate(row):
            field_name = _match_header(cell)
            if field_name and field_name not in mapping:
                mapping[field_name] = col_index
        if len(mapping) >= _MIN_MATCHED_HEADERS and "fecha" in mapping:
            return row_index, mapping
    return None


def ingest_spreadsheet(
    *, content: bytes, filename: str, sheet_name: str | None = None, forced_mapping: dict[str, int] | None = None, forced_header_row: int | None = None,
) -> SpreadsheetIngestResult:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheet_names = list(workbook.sheetnames)
    sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]

    all_rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    preview = SheetPreview(
        sheet_name=sheet.title,
        rows=[[("" if cell is None else str(cell)) for cell in row] for row in all_rows[:10]],
    )

    if forced_mapping is not None and forced_header_row is not None:
        header_row_index, mapping = forced_header_row, forced_mapping
        header_found = True
    else:
        detected = _find_header_row(all_rows)
        if detected is None:
            return SpreadsheetIngestResult(sheet_names, sheet.title, False, None, [], [], preview, [])
        header_row_index, mapping = detected
        header_found = True

    ready_rows: list[ReadyRow] = []
    ambiguous_rows: list[AmbiguousRow] = []
    fuel_charge_rows: list[FuelChargeRow] = []

    for offset, row in enumerate(all_rows[header_row_index + 1 :]):
        row_number = header_row_index + 2 + offset  # 1-indexed, human-facing
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        def cell(field_name: str) -> object:
            index = mapping.get(field_name)
            return row[index] if index is not None and index < len(row) else None

        operation_date = _parse_date(cell("fecha"))
        debe = _parse_amount(cell("debe"))
        haber = _parse_amount(cell("haber"))
        litres = _parse_amount(cell("litros"))
        currency_raw = str(cell("moneda") or "").strip().upper()
        currency = currency_raw if len(currency_raw) == 3 else ("ARS" if not currency_raw else None)
        description = str(cell("concepto") or "").strip() or _DEFAULT_DESCRIPTION
        reference_cell = cell("comprobante") or cell("referencia")
        reference = str(reference_cell).strip() if reference_cell else None

        extracted = {
            "operation_date": operation_date.isoformat() if operation_date else "",
            "debe": str(debe) if debe else "",
            "haber": str(haber) if haber else "",
            "currency": currency_raw,
            "description": description,
        }

        if operation_date is None:
            ambiguous_rows.append(AmbiguousRow(row_number, "No se pudo interpretar la fecha.", extracted))
            continue
        if currency is None:
            ambiguous_rows.append(AmbiguousRow(row_number, "La moneda no es clara (se esperaba un código de 3 letras, ej. ARS o USD).", extracted))
            continue
        if debe and haber:
            ambiguous_rows.append(AmbiguousRow(row_number, "La fila tiene importe en Debe y en Haber a la vez.", extracted))
            continue
        if not debe and not haber:
            if litres:
                fuel_charge_rows.append(
                    FuelChargeRow(
                        row_number=row_number,
                        operation_date=operation_date,
                        litres=litres,
                        description=description[:180],
                        currency=currency,
                        external_reference=reference,
                        row_hash_seed=f"{filename}:{sheet.title}:{row_number}:{operation_date}:{litres}:{currency}:fuel",
                    )
                )
                continue
            ambiguous_rows.append(AmbiguousRow(row_number, "No se encontró un importe en Debe ni en Haber.", extracted))
            continue

        amount, direction = (debe, "company") if debe else (haber, "counterparty")
        ready_rows.append(
            ReadyRow(
                row_number=row_number,
                operation_date=operation_date,
                description=description[:180],
                amount=amount,
                currency=currency,
                direction=direction,
                external_reference=reference,
                row_hash_seed=f"{filename}:{sheet.title}:{row_number}:{operation_date}:{amount}:{currency}:{direction}",
            )
        )

    return SpreadsheetIngestResult(sheet_names, sheet.title, header_found, header_row_index, ready_rows, ambiguous_rows, preview, fuel_charge_rows)
